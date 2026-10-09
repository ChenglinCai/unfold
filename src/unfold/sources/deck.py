"""PowerPoint decks: one anchor per slide, with the slide's speaker notes.

python-pptx reads only a slide's text runs, so equations would vanish. The reader
walks each paragraph itself, and turns the OMML of each equation into TeX.
"""

from pathlib import Path

from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]
from pptx import Presentation
from pptx.shapes.autoshape import Shape

from unfold.sources import Anchor, Meta, SourceDocument, build
from unfold.sources.omml import omml_tex
from unfold.sources.profile import quality

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _inline(node: etree._Element, parts: list[str]) -> None:
    """Append a paragraph's text and equations, in order."""
    for child in node:
        if child.tag == f"{A}t":
            parts.append(child.text or "")
        elif child.tag == f"{A}br":
            parts.append(" ")
        elif child.tag == f"{M}oMathPara":
            parts += [f" $${omml_tex(math)}$$ " for math in child.iter(f"{M}oMath")]
        elif child.tag == f"{M}oMath":
            parts.append(f"${omml_tex(child)}$")
        else:
            _inline(child, parts)


def shape_text(shape: Shape) -> str:
    """A shape's paragraphs, one per line, with each equation as TeX."""
    lines = []
    for paragraph in shape.element.iter(f"{A}p"):
        parts: list[str] = []
        _inline(paragraph, parts)
        lines.append(" ".join("".join(parts).split()))
    return "\n".join(lines).strip()


def read(path: Path, meta: Meta) -> SourceDocument:
    deck = Presentation(str(path))
    anchors: list[Anchor] = []
    words = 0
    for number, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        title_id = title_shape.shape_id if title_shape is not None else None
        parts = [
            shape_text(shape)
            for shape in slide.shapes
            if isinstance(shape, Shape)
            and shape.has_text_frame
            and shape.shape_id != title_id
        ]
        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = frame.text.strip() if frame is not None else ""
        words += sum(len(text.split()) for text in [title, *parts, notes])
        body = "\n\n".join(part for part in parts if part)
        if notes:
            body += f"\n\nSpeaker notes: {notes}"
        anchors.append(Anchor(f"s-{number}", "slide", title or f"Slide {number}", body))
    profile = {
        "format": "pptx",
        "size": {"slides": len(anchors), "words": words},
        "quality": quality(words, len(anchors), "slide"),
    }
    return build(meta, anchors, profile)
