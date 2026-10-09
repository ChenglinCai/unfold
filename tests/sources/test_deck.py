"""PowerPoint decks become one anchor per slide, with speaker notes."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.shapes.autoshape import Shape

from unfold.sources import Meta
from unfold.sources.deck import read


@pytest.fixture
def deck(tmp_path: Path) -> Path:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for title, body in [
        ("Variance", "Spread around the mean."),
        ("Bias", "Error from assumptions."),
    ]:
        slide = presentation.slides.add_slide(layout)
        heading, placeholder = slide.shapes.title, slide.placeholders[1]
        assert heading is not None
        assert isinstance(placeholder, Shape)
        heading.text = title
        placeholder.text_frame.text = body
    notes = presentation.slides[0].notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Square the distances first."
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


def test_each_slide_becomes_an_anchor(deck: Path) -> None:
    doc = read(deck, Meta(id="d", title="Deck", family="slides", origin="deck.pptx"))

    assert [(a.id, a.title) for a in doc.anchors] == [
        ("s-1", "Variance"),
        ("s-2", "Bias"),
    ]
    assert "Spread around the mean." in doc.anchors[0].text
    assert "Square the distances first." in doc.anchors[0].text
    assert doc.profile["size"] == {"slides": 2, "words": 13}
    assert doc.profile["needs"] == "fill-gaps"


def test_equations_on_a_slide_become_tex(tmp_path: Path) -> None:
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

    drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    math = (
        '<a14:m xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{}</a14:m>'
    )
    fraction = (
        "<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r>"
    )
    fraction += "</m:den></m:f>"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    heading, placeholder = slide.shapes.title, slide.placeholders[1]
    assert heading is not None and isinstance(placeholder, Shape)
    heading.text = "Rates"
    placeholder.text_frame.text = "The rate is "
    first = next(placeholder.element.iter(f"{drawing}p"))
    first.append(etree.fromstring(math.format(f"<m:oMath>{fraction}</m:oMath>")))
    placeholder.text_frame.add_paragraph()
    second = list(placeholder.element.iter(f"{drawing}p"))[1]
    display = "<m:oMathPara><m:oMath><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e>"
    display += "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath></m:oMathPara>"
    second.append(etree.fromstring(math.format(display)))
    path = tmp_path / "rates.pptx"
    presentation.save(str(path))

    doc = read(path, Meta(id="r", title="Rates", family="slides", origin="rates.pptx"))

    assert r"The rate is $\frac{a}{b}$" in doc.anchors[0].text
    assert "$${x}^{2}$$" in doc.anchors[0].text
